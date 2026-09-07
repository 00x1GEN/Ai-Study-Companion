from io import BytesIO
from uuid import uuid4
from docx import Document
from fastapi.testclient import TestClient
from app.main import app


def _docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Stack and Queue", level=1)
    doc.add_paragraph("A stack follows LIFO ordering. A queue follows FIFO ordering. Push and pop are stack operations. Enqueue and dequeue are queue operations.")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_full_reference_flow():
    email = f"qa-{uuid4().hex[:10]}@example.com"
    password = "password123"
    with TestClient(app) as client:
        # Register and obtain rotating token pair.
        r = client.post("/api/v1/auth/register", json={"name":"QA Student","email":email,"password":password})
        assert r.status_code == 201, r.text
        first = r.json()
        access = first["access_token"]
        old_refresh = first["refresh_token"]
        headers = {"Authorization": f"Bearer {access}"}

        # Refresh rotates; reusing the old token must fail.
        r = client.post("/api/v1/auth/refresh", json={"refresh_token":old_refresh})
        assert r.status_code == 200, r.text
        rotated = r.json()
        new_refresh = rotated["refresh_token"]
        access = rotated["access_token"]
        headers = {"Authorization": f"Bearer {access}"}
        assert client.post("/api/v1/auth/refresh", json={"refresh_token":old_refresh}).status_code == 401

        # Create a course.
        r = client.post("/api/v1/courses", headers=headers, json={"title":"QA Algorithms","description":"integration test"})
        assert r.status_code == 201, r.text
        course_id = r.json()["id"]

        # Upload DOCX and require successful ingestion into pgvector chunks.
        r = client.post(
            f"/api/v1/courses/{course_id}/materials",
            headers={"Authorization": f"Bearer {access}"},
            files={"file":("qa.docx", _docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        assert r.status_code == 201, r.text
        material = r.json()
        assert material["status"] == "ready", material
        assert material["chunk_count"] >= 1

        # Grounded RAG response must expose source metadata.
        r = client.post("/api/v1/ai/chat", headers=headers, json={"course_id":course_id,"message":"What ordering does a stack use?","provider":"mock"})
        assert r.status_code == 200, r.text
        assert r.json()["sources"], r.json()

        # AI learning-content generation.
        r = client.post("/api/v1/ai/flashcards/generate", headers=headers, json={"course_id":course_id,"count":2})
        assert r.status_code == 200, r.text
        assert r.json()["created"] >= 1
        cards = client.get(f"/api/v1/courses/{course_id}/flashcards", headers=headers).json()
        assert cards
        r = client.post(f"/api/v1/flashcards/{cards[0]['id']}/review", headers=headers, json={"known":True})
        assert r.status_code == 200, r.text

        r = client.post("/api/v1/ai/quiz/generate", headers=headers, json={"course_id":course_id,"count":2})
        assert r.status_code == 200, r.text
        quiz = client.get(f"/api/v1/courses/{course_id}/quiz", headers=headers)
        assert quiz.status_code == 200, quiz.text
        quiz_json = quiz.json()
        assert quiz_json["questions"]
        answers = {str(q["id"]): 0 for q in quiz_json["questions"]}
        result = client.post(f"/api/v1/quizzes/{quiz_json['id']}/submit", headers=headers, json={"answers":answers})
        assert result.status_code == 200, result.text

        # Progress and badge payload are structurally valid.
        progress = client.get("/api/v1/progress", headers=headers)
        assert progress.status_code == 200, progress.text
        assert progress.json()["xp"] >= 1
        assert isinstance(progress.json()["badges"], list)

        # Logout revokes the active refresh token.
        assert client.post("/api/v1/auth/logout", json={"refresh_token":new_refresh}).status_code == 200
        assert client.post("/api/v1/auth/refresh", json={"refresh_token":new_refresh}).status_code == 401
